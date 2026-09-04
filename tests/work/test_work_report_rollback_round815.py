from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook

from js_work.documents import WorkDocumentEngine
from js_work.routines.models import RoutineStatus, WorkRoutine
from js_work.routines.packing_details import PackingDetailsRoutineRunner
from js_work.routines.spreadsheet import WorkSpreadsheetRoutineRunner
from js_work.safe_output import publish_no_clobber, remove_published_link, staged_path


def _save_table_workbook(path: Path, *, sheet_name: str, rows: list[list[Any]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _spreadsheet_fixture(tmp_path: Path) -> tuple[WorkRoutine, Path]:
    _save_table_workbook(
        tmp_path / "source.xlsx",
        sheet_name="Source",
        rows=[["name", "qty"], ["authorized", 1]],
    )
    _save_table_workbook(
        tmp_path / "template.xlsx",
        sheet_name="Result",
        rows=[["name", "qty"], [None, None]],
    )
    routine = WorkRoutine(
        routine_id="rollback-probe",
        name="Rollback probe",
        trigger_phrases=["rollback probe"],
        routine_type="spreadsheet_template",
        status=RoutineStatus.ENABLED,
        field_mapping={"name": "name", "qty": "qty"},
        validation_rules={"required_fields": ["name", "qty"]},
    )
    return routine, tmp_path / "reports" / "result.xlsx"


def _swap_published_parent(
    parent: Path,
    parked: Path,
    replacement: Path,
    artifact_name: str,
    *,
    strategy: str,
) -> None:
    parent.rename(parked)
    if strategy == "symlink":
        replacement.mkdir()
        parent.symlink_to(replacement, target_is_directory=True)
        return
    if strategy == "hardlink":
        parent.mkdir()
        os.link(parked / artifact_name, parent / artifact_name)
        return
    raise AssertionError(f"unsupported swap strategy: {strategy}")


def test_descriptor_rollback_reports_identity_uncertainty_without_unlinking_replacement(
    tmp_path: Path,
) -> None:
    output = tmp_path / "reports" / "result.xlsx"
    with staged_path(output) as staged:
        staged.write_bytes(b"authorized artifact")
        publish_no_clobber(staged, output, "output already exists")
        output.unlink()
        output.write_bytes(b"replacement artifact")

        with pytest.raises(RuntimeError, match="rollback could not be confirmed"):
            remove_published_link(staged, output)

        assert output.read_bytes() == b"replacement artifact"


@pytest.mark.parametrize("strategy", ["symlink", "hardlink"])
def test_spreadsheet_report_failure_rolls_back_through_original_parent_anchor(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    strategy: str,
) -> None:
    import js_work.routines.spreadsheet as spreadsheet_module

    routine, output = _spreadsheet_fixture(tmp_path)
    parent = output.parent
    parked = tmp_path / "reports-parked"
    replacement = tmp_path / "replacement"
    swapped = False

    def fail_report_write(
        _path: Path,
        _payload: dict[str, Any],
        _message: str,
        **_kwargs: Any,
    ) -> None:
        nonlocal swapped
        _swap_published_parent(
            parent,
            parked,
            replacement,
            output.name,
            strategy=strategy,
        )
        swapped = True
        raise OSError("synthetic validation report write failure")

    monkeypatch.setattr(spreadsheet_module, "write_json_no_clobber", fail_report_write)

    with pytest.raises(OSError, match="synthetic validation report write failure"):
        WorkSpreadsheetRoutineRunner(tmp_path, tmp_path / "state").run(
            routine=routine,
            source_path="source.xlsx",
            template_path="template.xlsx",
            output_path="reports/result.xlsx",
        )

    assert swapped is True
    assert not (parked / output.name).exists()
    assert not (parked / "result.validation.json").exists()
    if strategy == "symlink":
        assert list(replacement.iterdir()) == []
    else:
        assert output.is_file(), "rollback must not unlink a replacement-directory hardlink"


@pytest.mark.parametrize("operation", ["create", "replace"])
@pytest.mark.parametrize("strategy", ["symlink", "hardlink"])
def test_word_report_failure_rolls_back_create_and_replace_through_original_anchor(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    operation: str,
    strategy: str,
) -> None:
    engine = WorkDocumentEngine(tmp_path)
    output = tmp_path / "reports" / f"{operation}.docx"
    parent = output.parent
    parked = tmp_path / f"reports-{operation}-parked"
    replacement = tmp_path / f"replacement-{operation}"
    if operation == "replace":
        source = Document()
        source.add_paragraph("authorized source text")
        source.save(tmp_path / "source.docx")
    swapped = False

    def fail_report_write(
        _cls: type[WorkDocumentEngine],
        _path: Path,
        _payload: dict[str, Any],
        **_kwargs: Any,
    ) -> None:
        nonlocal swapped
        _swap_published_parent(
            parent,
            parked,
            replacement,
            output.name,
            strategy=strategy,
        )
        swapped = True
        raise OSError("synthetic Word validation report write failure")

    monkeypatch.setattr(
        WorkDocumentEngine,
        "_write_json_no_clobber",
        classmethod(fail_report_write),
    )

    with pytest.raises(OSError, match="synthetic Word validation report write failure"):
        if operation == "create":
            engine.create_word(
                output_path=f"reports/{operation}.docx",
                title="Synthetic",
                sections=[],
            )
        else:
            engine.replace_word(
                source_path="source.docx",
                output_path=f"reports/{operation}.docx",
                replacements={"authorized source text": "approved replacement text"},
            )

    assert swapped is True
    assert not (parked / output.name).exists()
    assert not (parked / f"{operation}.validation.json").exists()
    assert list(parked.glob(f".{operation}.*.docx")) == []
    if strategy == "symlink":
        assert list(replacement.iterdir()) == []
    else:
        assert output.is_file(), "rollback must not unlink a replacement-directory hardlink"


def test_packing_report_commit_rejects_replaced_parent_and_rolls_back(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import js_work.routines.packing_details as packing_module
    from js_work.safe_output import write_json_no_clobber as real_write_json

    _save_table_workbook(
        tmp_path / "source.xlsx",
        sheet_name="Source",
        rows=[
            ["FABRICS", "PON.", "COLOR", "ROLL NO", "QTY(M)", "ROLL NO", "QTY(M)"],
            ["A", "P1", "WHITE", 1, 10, None, None],
            ["TOTAL", None, None, "1 roll", "=SUM(E2:E2)", None, None],
        ],
    )
    _save_table_workbook(
        tmp_path / "template.xlsx",
        sheet_name="PACKING DETAILS",
        rows=[
            ["FABRICS", "PON.", "COLOR", "ROLL NO", "QTY(M)", "ROLL NO", "QTY(M)"],
            [None, None, None, None, None, None, None],
            ["TOTAL", None, None, 1, "=SUM(E2:E2)+SUM(G2:G2)", None, None],
        ],
    )
    output = tmp_path / "reports" / "packing.xlsx"
    parked = tmp_path / "reports-parked"
    swapped = False

    def swap_then_write_report(
        path: Path,
        payload: dict[str, Any],
        message: str,
        **kwargs: Any,
    ) -> None:
        nonlocal swapped
        output.parent.rename(parked)
        output.parent.mkdir()
        swapped = True
        real_write_json(path, payload, message, **kwargs)

    monkeypatch.setattr(packing_module, "write_json_no_clobber", swap_then_write_report)

    with pytest.raises((OSError, ValueError)):
        PackingDetailsRoutineRunner(tmp_path).run(
            source_path="source.xlsx",
            template_path="template.xlsx",
            output_path="reports/packing.xlsx",
        )

    assert swapped is True
    assert not (parked / output.name).exists()
    assert not (parked / "packing.validation.json").exists()
    assert not output.exists()
    assert not output.with_suffix(".validation.json").exists()
